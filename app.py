"""RevenueLive: channel-scoped revenue reporting and audited spreadsheet imports."""
import argparse
import csv
import datetime as dt
import functools
import hashlib
import io
import itertools
import json
import os
from pathlib import Path
import secrets
import logging
import sqlite3
import threading
import time
from contextlib import closing
from decimal import Decimal, InvalidOperation

from flask import Flask, Request, g, jsonify, request, send_from_directory, redirect
from werkzeug.security import check_password_hash, generate_password_hash
from werkzeug.utils import secure_filename
from werkzeug.exceptions import HTTPException
from insight_presets import INSIGHT_PRESETS
from config import Settings, load_environment
from database import Database, INTEGRITY_ERRORS

ROOT = Path(__file__).resolve().parent
HEADERS = ['Date', 'Channel Name', 'Views', 'Ad Impressions', 'Ad Revenue', 'Sponsorship/Others', 'Total Revenue']
AUDIT_GENESIS = '0' * 64


def audit_digest(row):
    payload={key:row[key] for key in ('id','created','user_id','actor','action','detail','prev_hash')}
    return hashlib.sha256(json.dumps(payload,sort_keys=True,separators=(',',':'),ensure_ascii=True).encode()).hexdigest()


def verify_audit_chain(connection):
    previous=AUDIT_GENESIS;count=0
    for row in connection.execute('SELECT * FROM audit ORDER BY id'):
        if row['prev_hash']!=previous or row['entry_hash']!=audit_digest(row):
            return False,previous,count,row['id']
        previous=row['entry_hash'];count+=1
    return True,previous,count,None


class InvalidData(ValueError):
    pass


class MemoryUploadRequest(Request):
    def _get_file_stream(self, total_content_length, content_type, filename=None, content_length=None):
        # Request size is bounded by MAX_CONTENT_LENGTH; avoid multipart disk spooling.
        return io.BytesIO()


def parse_upload(content, suffix):
    if suffix == '.xls':
        import xlrd
        book = xlrd.open_workbook(file_contents=content)
        sheet = book.sheet_by_index(0)
        if sheet.nrows > 20001:
            raise InvalidData('Maximum 20,000 data rows per upload.')
        rows = []
        for i in range(sheet.nrows):
            row = sheet.row_values(i)
            if i and sheet.cell_type(i, 0) == xlrd.XL_CELL_DATE:
                row[0] = xlrd.xldate_as_datetime(row[0], book.datemode)
            rows.append(row)
    elif suffix == '.xlsx':
        from openpyxl import load_workbook
        import zipfile
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            if sum(x.file_size for x in archive.infolist()) > 100_000_000:
                raise InvalidData('Workbook expands beyond the 100 MB limit.')
        book = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        try:
            rows = list(itertools.islice(book.worksheets[0].iter_rows(values_only=True), 20002))
        finally:
            book.close()
    elif suffix == '.csv':
        rows = list(itertools.islice(csv.reader(io.StringIO(content.decode('utf-8-sig'))), 20002))
    else:
        raise InvalidData('Upload an XLS, XLSX or UTF-8 CSV file.')
    if not rows or [str(x or '').strip() for x in rows[0]] != HEADERS:
        raise InvalidData('Columns must match: ' + ', '.join(HEADERS))
    if len(rows) > 20001:
        raise InvalidData('Maximum 20,000 data rows per upload.')
    result, seen = [], set()
    for number, row in enumerate(rows[1:], 2):
        if all(x is None or str(x).strip() == '' for x in row):
            continue
        if len(row) != 7:
            raise InvalidData(f'Row {number}: expected seven columns.')
        date = row[0]
        if isinstance(date, dt.datetime):
            date = date.date()
        elif not isinstance(date, dt.date):
            try:
                date = dt.date.fromisoformat(str(date).strip())
            except ValueError:
                raise InvalidData(f'Row {number}: use an Excel date or YYYY-MM-DD.')
        channel = str(row[1] or '').strip()
        if not channel or len(channel) > 120:
            raise InvalidData(f'Row {number}: invalid channel name.')
        key = (date.isoformat(), channel.casefold())
        if key in seen:
            raise InvalidData(f'Row {number}: duplicate date/channel inside the file.')
        seen.add(key)
        values = []
        for index, value in enumerate(row[2:], 2):
            try:
                numeric = Decimal(str(value).strip())
                if not numeric.is_finite() or numeric < 0 or numeric > Decimal('1000000000000'):
                    raise ValueError()
                scaled = numeric * (100 if index >= 4 else 1)
                if scaled != scaled.to_integral_value():
                    raise ValueError()
                values.append(int(scaled))
            except (InvalidOperation, ValueError):
                raise InvalidData(f'Row {number}: {HEADERS[index]} must be non-negative; counts are integers and INR allows two decimals.')
        if values[2] + values[3] != values[4]:
            raise InvalidData(f'Row {number}: total revenue must equal ad revenue plus sponsorship/others.')
        result.append(dict(day=date.isoformat(), channel=channel, views=values[0], impressions=values[1], ad=values[2], other=values[3], total=values[4]))
    if not result:
        raise InvalidData('No data rows found.')
    return result


def create_app(data_dir=None):
    settings = Settings.from_env(data_dir)
    public_origin = settings.app_url
    app = Flask(__name__, static_folder='static')
    app.request_class = MemoryUploadRequest
    data = settings.data_dir
    data.mkdir(parents=True, exist_ok=True)
    uploads_dir = settings.upload_dir
    app.config.update(DATA_DIR=data, UPLOAD_DIR=uploads_dir, SETTINGS=settings,
                      MAX_CONTENT_LENGTH=10*1024*1024, TRUSTED_HOSTS=list(settings.hosts) or None)
    cookie_name = settings.cookie_name
    database = Database(settings)
    app.extensions['database'] = database

    def db():
        if 'db' not in g:
            g.db = database.connect()
        return g.db

    @app.teardown_appcontext
    def close(_error):
        connection = g.pop('db', None)
        if connection:
            connection.close()

    with app.app_context():
        if database.mysql:
            database.check_schema()
            valid, _, _, broken = verify_audit_chain(db())
            if not valid:
                raise RuntimeError(f'Audit chain verification failed at event {broken}.')
        else:
            from sqlite_legacy import initialize
            initialize(db, audit_digest, verify_audit_chain, AUDIT_GENESIS)
    def log(action, detail='', user=None):
        db().begin_write()
        user=user or g.user
        valid,_,_,broken=verify_audit_chain(db())
        if not valid:
            raise RuntimeError(f'Audit chain verification failed at event {broken}.')
        last=db().execute('SELECT id,entry_hash FROM audit ORDER BY id DESC LIMIT 1').fetchone()
        event=dict(id=last['id']+1 if last else 1,created=dt.datetime.now(dt.timezone.utc).isoformat(),
                   user_id=user['id'],actor=user['username'],action=action,detail=str(detail),
                   prev_hash=last['entry_hash'] if last else AUDIT_GENESIS)
        event['entry_hash']=audit_digest(event)
        db().execute('INSERT INTO audit(id,created,user_id,action,detail,actor,prev_hash,entry_hash) VALUES (:id,:created,:user_id,:action,:detail,:actor,:prev_hash,:entry_hash)',event)

    def permitted():
        if g.user['role'] == 'admin':
            return db().execute('SELECT * FROM channels WHERE id NOT IN (SELECT channel_id FROM archived_channels) ORDER BY LOWER(name)').fetchall()
        return db().execute('SELECT c.* FROM channels c JOIN assignments a ON c.id=a.channel_id WHERE a.user_id=? AND c.id NOT IN (SELECT channel_id FROM archived_channels) ORDER BY LOWER(c.name)', (g.user['id'],)).fetchall()

    def require(*roles):
        def decorator(fn):
            @functools.wraps(fn)
            def wrapped(*args, **kwargs):
                if not getattr(g, 'user', None):
                    return jsonify(error='Please sign in.'), 401
                if g.user['must_change'] and request.path not in {'/api/me','/api/password','/api/logout'}:
                    return jsonify(error='Change your temporary password first.'), 403
                if roles and g.user['role'] not in roles:
                    return jsonify(error='You do not have permission for this action.'), 403
                if roles == ('admin',) and not settings.admin_host(request.host):
                    return jsonify(error='Use the admin portal for this action.'), 403
                return fn(*args, **kwargs)
            return wrapped
        return decorator

    @app.before_request
    def auth():
        protected_page = request.path in {'/admin', '/user'}
        if public_origin and not settings.request_origin(request.host):
            return jsonify(error='Unknown application host.'), 400
        if not request.path.startswith('/api/') and not protected_page:
            return
        g.user = None
        if database.mysql and request.method in {'POST', 'PUT', 'PATCH', 'DELETE'}:
            db().begin_write()
        raw = request.cookies.get(cookie_name, '')
        digest = hashlib.sha256(raw.encode()).hexdigest()
        session = db().execute('SELECT s.*,u.username,u.role,u.active,u.must_change FROM sessions s JOIN users u ON u.id=s.user_id WHERE token=? AND expires>? AND active=1', (digest,time.time())).fetchone()
        if session:
            g.user = dict(id=session['user_id'], username=session['username'], role=session['role'], must_change=session['must_change'])
            g.user['super_admin'] = bool(db().execute('SELECT 1 FROM super_admin WHERE user_id=?',(session['user_id'],)).fetchone())
            g.session = session
            if settings.portal_mode == 'split' and (g.user['role'] == 'admin') != settings.admin_host(request.host):
                return jsonify(error='Sign in through the correct portal.'), 403
        if protected_page:
            if not g.user:
                return redirect('/login')
            if request.path == '/admin' and (g.user['role'] != 'admin' or not settings.admin_host(request.host)):
                return 'Admin access required.', 403
        if request.method in {'POST','PUT','PATCH','DELETE'}:
            # Reject cross-site writes even on login, where no session exists yet.
            origin = request.headers.get('Origin')
            if origin and origin != (settings.request_origin(request.host) or request.host_url.rstrip('/')):
                return jsonify(error='Cross-origin requests are not allowed.'), 403
            if request.path not in {'/api/login','/api/account/request','/api/account/complete'} and (not session or not secrets.compare_digest(request.headers.get('X-CSRF-Token',''),session['csrf'])):
                return jsonify(error='Session expired. Sign in again.'), 403

    @app.after_request
    def headers(response):
        response.headers['X-Content-Type-Options']='nosniff'
        response.headers['X-Frame-Options']='DENY'
        response.headers['Referrer-Policy']='same-origin'
        response.headers['Content-Security-Policy']="default-src 'self'; script-src 'self'; style-src 'self'; img-src 'self' data:; connect-src 'self'; frame-ancestors 'none'; base-uri 'self'; form-action 'self'"
        if public_origin:
            response.headers['Strict-Transport-Security']='max-age=31536000'
        if request.path.startswith('/api/') or request.path in {'/', '/login', '/admin', '/user'} or request.path.startswith('/static/'):
            response.headers['Cache-Control']='no-store'
        return response

    @app.errorhandler(413)
    def large(_):
        return jsonify(error='Maximum file size is 10 MB.'),413

    @app.errorhandler(InvalidData)
    def invalid(error):
        return jsonify(error=str(error)),400

    @app.errorhandler(HTTPException)
    def http_error(error):
        if request.path.startswith('/api/'):
            return jsonify(error='This server needs an application update.' if error.code==404 else error.description),error.code
        return error

    @app.get('/')
    @app.get('/login')
    @app.get('/admin')
    @app.get('/user')
    def index():
        return send_from_directory(ROOT / 'static','index.html')

    @app.get('/health')
    def health():
        return jsonify(service='revenuelive', status='ok', uploads_version=2)

    @app.post('/api/login')
    def login():
        value = request.get_json(silent=True) or {}
        ip = request.remote_addr
        attempt = db().execute('SELECT * FROM attempts WHERE ip=?',(ip,)).fetchone()
        if attempt and attempt['expires']>time.time() and attempt['failures']>=10:
            return jsonify(error='Too many attempts. Try again in 15 minutes.'),429
        user = db().execute('SELECT * FROM users WHERE username=? AND active=1',(str(value.get('username','')).strip()[:100],)).fetchone()
        password = str(value.get('password',''))
        if not user or len(password)>256 or not check_password_hash(user['password'], password):
            failures = attempt['failures']+1 if attempt and attempt['expires']>time.time() else 1
            db().upsert('attempts',dict(ip=ip,failures=failures,expires=time.time()+900))
            db().commit()
            return jsonify(error='Invalid username or password.'),401
        raw, csrf = secrets.token_urlsafe(32), secrets.token_urlsafe(32)
        db().execute('DELETE FROM attempts WHERE ip=?',(ip,))
        db().execute('DELETE FROM sessions WHERE expires<?',(time.time(),))
        if user['role'] == 'admin' and not settings.admin_host(request.host):
            return jsonify(error='Sign in through the admin portal.'), 403
        if settings.portal_mode == 'split' and settings.admin_host(request.host) and user['role'] != 'admin':
            return jsonify(error='Sign in through the user portal.'), 403
        db().execute('INSERT INTO sessions VALUES (?,?,?,?)',(hashlib.sha256(raw.encode()).hexdigest(),user['id'],csrf,time.time()+settings.session_ttl))
        log('user_signed_in',user=user)
        db().commit()
        response=jsonify(ok=True)
        response.set_cookie(cookie_name,raw,httponly=True,samesite=settings.cookie_samesite,
                            secure=settings.cookie_secure,max_age=settings.session_ttl)
        return response

    @app.get('/api/me')
    @require()
    def me():
        return jsonify(user=g.user,csrf=g.session['csrf'],channels=[dict(c) for c in permitted()],
                       home='/admin' if g.user['role']=='admin' else '/user')

    @app.post('/api/logout')
    @require()
    def logout():
        db().execute('DELETE FROM sessions WHERE token=?',(g.session['token'],))
        log('user_signed_out');db().commit()
        response=jsonify(ok=True);response.delete_cookie(cookie_name,secure=settings.cookie_secure,
                                                        httponly=True,samesite=settings.cookie_samesite);return response

    @app.post('/api/password')
    @require()
    def password():
        body=request.get_json() or {}
        value=str(body.get('password',''))
        if not 12<=len(value)<=256:
            raise InvalidData('Password must be 12 to 256 characters.')
        current=db().execute('SELECT password FROM users WHERE id=?',(g.user['id'],)).fetchone()[0]
        if not check_password_hash(current,str(body.get('current',''))):
            raise InvalidData('Current password is incorrect.')
        db().execute('UPDATE users SET password=?,must_change=0 WHERE id=?',(generate_password_hash(value),g.user['id']))
        db().execute('DELETE FROM email_tokens WHERE user_id=?',(g.user['id'],))
        db().execute('DELETE FROM sessions WHERE user_id=? AND token<>?',(g.user['id'],g.session['token']))
        log('password_changed');db().commit();return jsonify(ok=True)

    def report_rows(available_dates=False):
        ids={c['id'] for c in permitted()}
        selected=[v for v in request.args.getlist('channel') if v]
        if selected == ['none']:
            ids=set()
        elif selected:
            try:
                chosen={int(value) for value in selected}
            except ValueError:
                raise InvalidData('Invalid channel.')
            if not chosen <= ids:
                raise InvalidData('Channel is not assigned to your account.')
            ids=chosen
        for key in ('start','end'):
            if request.args.get(key):
                try:
                    dt.date.fromisoformat(request.args[key])
                except ValueError:
                    raise InvalidData('Use YYYY-MM-DD dates.')
        start,end=request.args.get('start',''),request.args.get('end','9999-12-31') or '9999-12-31'
        if start>end:
            start,end=end,start
        placeholders=','.join('?' for _ in ids) or 'NULL'
        visible="upload_id NOT IN (SELECT id FROM uploads WHERE archived=1 OR state!='committed' OR file_deleted=1)"
        if available_dates:
            return [r[0] for r in db().execute(f'SELECT DISTINCT day FROM records WHERE channel_id IN ({placeholders}) AND day NOT IN (SELECT day FROM hidden_dates) AND {visible} ORDER BY day',tuple(sorted(ids)))]
        return [dict(r) for r in db().execute(f'SELECT r.*,c.name AS channel FROM records r JOIN channels c ON c.id=r.channel_id WHERE r.channel_id IN ({placeholders}) AND day>=? AND day<=? AND day NOT IN (SELECT day FROM hidden_dates) AND {visible} ORDER BY day DESC,LOWER(c.name)',(*sorted(ids),start,end))]

    @app.get('/api/admin/dates')
    @require('admin')
    def reporting_dates():
        counts={r['day']:r['channels'] for r in db().execute('SELECT day,COUNT(*) AS channels FROM records GROUP BY day')}
        hidden={r['day'] for r in db().execute('SELECT day FROM hidden_dates')}
        return jsonify(rows=[dict(day=day,channels=counts.get(day,0),hidden=day in hidden) for day in sorted(counts.keys()|hidden,reverse=True)])

    @app.post('/api/admin/dates/<day>/visibility')
    @require('admin')
    def set_date_visibility(day):
        try:
            dt.date.fromisoformat(day)
        except ValueError:
            raise InvalidData('Use a YYYY-MM-DD date.')
        hidden=(request.get_json(silent=True) or {}).get('hidden')
        if type(hidden) is not bool:
            raise InvalidData('Choose whether this date is hidden.')
        db().begin_write()
        if hidden:
            if not db().execute('SELECT 1 FROM records WHERE day=?',(day,)).fetchone():
                raise InvalidData('No published records exist for this date.')
            db().upsert('hidden_dates',dict(day=day,user_id=g.user['id'],created=dt.datetime.now(dt.timezone.utc).isoformat()),ignore=True)
        else:
            db().execute('DELETE FROM hidden_dates WHERE day=?',(day,))
        log('date_hidden' if hidden else 'date_restored',day)
        db().commit();return jsonify(ok=True)

    @app.get('/api/admin/audit')
    @require('admin')
    def audit_history():
        valid,previous,count,broken=verify_audit_chain(db())
        events=[dict(r) for r in db().execute('SELECT id,created,user_id,actor,action,detail,prev_hash,entry_hash FROM audit ORDER BY id DESC LIMIT 100')]
        return jsonify(events=events,valid=valid,count=count,head=previous,broken_event=broken)

    @app.get('/api/admin/audit/export')
    @require('admin')
    def export_audit():
        events=[dict(r) for r in db().execute('SELECT id,created,user_id,actor,action,detail,prev_hash,entry_hash FROM audit ORDER BY id')]
        valid,head,_,broken=verify_audit_chain(db())
        response=jsonify(events=events,head=head,valid=valid,broken_event=broken)
        response.headers['Content-Disposition']='attachment; filename=revenue-audit-chain.json'
        return response

    @app.get('/api/report')
    @require()
    def report():
        rows=report_rows()
        totals={key:sum(r[key] for r in rows) for key in ('views','impressions','ad','other','total')}
        return jsonify(rows=rows,totals=totals,currency='INR',money_unit='paise',available_dates=report_rows(available_dates=True))

    @app.get('/api/graph-presets')
    @require()
    def graph_presets():
        return jsonify(rows=[dict(id=r['id'],name=r['name'],config=json.loads(r['config'])) for r in db().execute('SELECT * FROM graph_presets WHERE user_id=? ORDER BY name',(g.user['id'],))],examples=INSIGHT_PRESETS)

    @app.post('/api/graph-presets')
    @require()
    def save_graph_preset():
        body=request.get_json(silent=True) or {}
        name=body.get('name','')
        config=body.get('config')
        if not isinstance(name,str) or not 1<=len(name.strip())<=80 or not isinstance(config,dict):
            return jsonify(error='Enter a preset name (1-80 characters).'),400
        allowed={'views','impressions','ad','other','total'}
        if any(not isinstance(config.get(key),str) for key in ('type','group','first','second')) or config.get('type') not in {'bar','line','grouped','mixed','pie'} or config.get('group') not in {'day','channel','leader'} or config.get('first') not in allowed or config.get('second') not in allowed|{'none'}:
            return jsonify(error='Invalid chart configuration.'),400
        clean={k:config[k] for k in ('type','group','first','second')}
        for key in ('start','end'):
            value=config.get(key,'')
            if not isinstance(value,str):
                return jsonify(error='Invalid preset date.'),400
            try:
                if value:
                    dt.date.fromisoformat(value)
            except (TypeError,ValueError):
                return jsonify(error='Invalid preset date.'),400
            clean[key]=value
        if clean['start'] and clean['end'] and clean['start']>clean['end']:
            clean['start'],clean['end']=clean['end'],clean['start']
        if 'channels' in config:
            channels=config['channels']
            permitted_ids={c['id'] for c in permitted()}
            if not isinstance(channels,list) or any(type(c) is not int or c not in permitted_ids for c in channels):
                return jsonify(error='Invalid channel selection.'),400
            clean['channels']=sorted(set(channels))
        try:
            db().execute('INSERT INTO graph_presets(user_id,name,config) VALUES (?,?,?)',(g.user['id'],name.strip(),json.dumps(clean)))
            log('graph_preset_saved',name.strip())
            db().commit()
        except INTEGRITY_ERRORS:
            return jsonify(error='A preset with that name already exists. Choose another name.'),409
        return jsonify(ok=True)

    @app.delete('/api/graph-presets/<int:pid>')
    @require()
    def delete_graph_preset(pid):
        result=db().execute('DELETE FROM graph_presets WHERE id=? AND user_id=?',(pid,g.user['id']))
        if not result.rowcount:
            return jsonify(error='Preset not found.'),404
        log('graph_preset_deleted',str(pid))
        db().commit()
        return jsonify(ok=True)

    @app.get('/api/export')
    @require()
    def download():
        buffer=io.StringIO();writer=csv.writer(buffer);writer.writerow(HEADERS)
        for r in report_rows():
            channel=r['channel']
            if channel.startswith(('=','+','-','@')):
                channel="'"+channel
            writer.writerow([r['day'],channel,r['views'],r['impressions'],*[f'{r[k]/100:.2f}' for k in ('ad','other','total')]])
        response=app.response_class('\ufeff'+buffer.getvalue(),mimetype='text/csv')
        response.headers['Content-Disposition']='attachment; filename=revenue.csv'
        return response

    def resolve_rows(rows):
        allowed={c['id'] for c in permitted()}
        names={c['name'].casefold():c['id'] for c in db().execute('SELECT * FROM channels')}
        resolved=[]
        for row in rows:
            cid=names.get(row['channel'].casefold())
            if cid is None or cid not in allowed:
                raise InvalidData('File contains unknown or unassigned channels. Ask the admin to register/assign all channels before uploading.')
            resolved.append({**row,'channel_id':cid})
        return resolved

    def row_values(row):
        return tuple(row[key] for key in ('views','impressions','ad','other','total'))

    def row_fingerprint(rows):
        return hashlib.sha256(json.dumps(sorted((r['day'],r['channel_id'],*row_values(r)) for r in rows),separators=(',',':')).encode()).digest()

    @app.post('/api/uploads/preview')
    @require('admin','uploader')
    def preview():
        incoming=request.files.get('file')
        if not incoming or not incoming.filename:
            raise InvalidData('Choose a file.')
        content=incoming.read()
        try:
            rows=parse_upload(content,Path(incoming.filename).suffix.lower())
        except InvalidData:
            raise
        except Exception:
            raise InvalidData('Unable to read workbook. Check its format and date cells.')
        rows=resolve_rows(rows)
        digest=hashlib.sha256(content).hexdigest()
        db().begin_write()
        fingerprint=row_fingerprint(rows)
        for existing in db().execute("SELECT rows_json FROM uploads WHERE state='pending'"):
            if row_fingerprint(json.loads(existing['rows_json']))==fingerprint:
                raise InvalidData('The same data is already awaiting publication. Publish or discard that preview first.')
        originals=[]
        for row in rows:
            old=db().execute('SELECT * FROM records WHERE day=? AND channel_id=?',(row['day'],row['channel_id'])).fetchone()
            originals.append(dict(old) if old else None)
        unchanged=sum(old is not None and row_values(old)==row_values(row) for row,old in zip(rows,originals))
        if unchanged==len(rows):
            raise InvalidData('All rows already match published data. No changes to publish.')
        uid=secrets.token_hex(16)
        filename=secure_filename(incoming.filename) or 'upload'
        try:
            db().execute('INSERT INTO uploads(id,user_id,filename,digest,created,state,rows_json,preview_json) VALUES (?,?,?,?,?,?,?,?)',(uid,g.user['id'],filename,digest,dt.datetime.now(dt.timezone.utc).isoformat(),'pending',json.dumps(rows),json.dumps(originals)))
            log('upload_submitted',json.dumps(dict(id=uid,filename=filename,sha256=digest,rows=len(rows)),separators=(',',':')))
            db().commit()
        except Exception:
            db().rollback()
            raise
        return jsonify(id=uid,rows=rows,duplicates=sum(old is not None and row_values(old)!=row_values(row) for row,old in zip(rows,originals)),unchanged=unchanged)

    @app.post('/api/uploads/<uid>/commit')
    @require('admin','uploader')
    def commit(uid):
        db().begin_write()
        upload=db().execute("SELECT * FROM uploads WHERE id=? AND (user_id=? OR ?='admin')",(uid,g.user['id'],g.user['role'])).fetchone()
        if not upload or upload['state']!='pending':
            raise InvalidData('Upload is unavailable or already published.')
        publish_preview(upload,(request.get_json(silent=True) or {}).get('replace') is True)
        db().commit();return jsonify(ok=True)

    def publish_preview(upload,replace):
        uid=upload['id']
        rows=resolve_rows(json.loads(upload['rows_json']))
        original=json.loads(upload['preview_json'])
        if any(old is not None and row_values(old)!=row_values(row) for row,old in zip(rows,original)) and not replace:
            raise InvalidData('Confirm replacement of existing date/channel records.')
        changed=0
        for row,old in zip(rows,original):
            current=db().execute('SELECT * FROM records WHERE day=? AND channel_id=?',(row['day'],row['channel_id'])).fetchone()
            if (dict(current) if current else None)!=old:
                raise InvalidData('Data changed after preview. Upload again to review the latest version.')
            if old is not None and row_values(old)==row_values(row):
                continue
            db().execute('INSERT INTO revisions VALUES (?,?,?,?)',(uid,row['day'],row['channel_id'],json.dumps(old)))
            db().upsert('records',{key:row[key] for key in ('day','channel_id','views','impressions','ad','other','total')} | {'upload_id':uid})
            changed+=1
        if not changed:
            raise InvalidData('All rows already match published data. No changes to publish.')
        db().execute("UPDATE uploads SET state='committed',archived=0 WHERE id=?",(uid,))
        log('upload_published',json.dumps(dict(id=uid,filename=upload['filename'],changed=changed),separators=(',',':')))

    @app.get('/api/uploads')
    @require('admin','uploader')
    def uploads():
        show_archived=request.args.get('show_archived')=='1'
        conditions=["u.state!='deleted'"]
        if not show_archived:
            conditions.append('u.archived=0')
        args=[]
        if g.user['role']!='admin':
            conditions.append('u.user_id=?');args.append(g.user['id'])
        where=' WHERE '+' AND '.join(conditions) if conditions else ''
        rows=db().execute('''SELECT u.id,u.filename,u.created,u.state,u.archived,u.file_deleted,u.rows_json,u.preview_json,v.username,
            (SELECT COUNT(*) FROM revisions WHERE upload_id=u.id) AS changed_rows,
            (SELECT COUNT(*) FROM records WHERE upload_id=u.id) AS live_rows,
            (SELECT COUNT(*) FROM records r WHERE r.upload_id=u.id AND r.day NOT IN (SELECT day FROM hidden_dates)
                AND r.channel_id NOT IN (SELECT channel_id FROM archived_channels)
                AND u.archived=0 AND u.state='committed' AND u.file_deleted=0) AS visible_rows
            FROM uploads u JOIN users v ON v.id=u.user_id'''+where+' ORDER BY u.created DESC LIMIT 100',args).fetchall()
        result=[]
        for row in rows:
            item={key:row[key] for key in row.keys() if key not in ('rows_json','preview_json')}
            original=json.loads(row['preview_json'])
            incoming=json.loads(row['rows_json'])
            item['total_rows']=len(incoming)
            days=sorted({r['day'] for r in incoming})
            item['start']=days[0] if days else None
            item['end']=days[-1] if days else None
            item['channel_count']=len({r['channel_id'] for r in incoming})
            item['replacements']=sum(old is not None and row_values(old)!=row_values(new) for old,new in zip(original,incoming)) if row['state'] in ('pending','rejected') else 0
            result.append(item)
        return jsonify(rows=result,uploads_version=2)

    @app.get('/api/uploads/<uid>/file')
    @require('admin','uploader')
    def upload_file(uid):
        upload=db().execute('SELECT * FROM uploads WHERE id=?',(uid,)).fetchone()
        if not upload or (g.user['role']!='admin' and upload['user_id']!=g.user['id']):
            return jsonify(error='Upload not found.'),404
        if upload['file_deleted'] or upload['state']=='deleted':
            return jsonify(error='Upload data is no longer available for download.'),404
        buffer=io.StringIO();writer=csv.writer(buffer);writer.writerow(HEADERS)
        for row in json.loads(upload['rows_json']):
            channel=row['channel']
            if channel.startswith(('=','+','-','@')):
                channel="'"+channel
            writer.writerow([row['day'],channel,row['views'],row['impressions'],
                             *[format(Decimal(row[key])/100,'.2f') for key in ('ad','other','total')]])
        filename=Path(secure_filename(upload['filename']) or 'upload').stem+'.csv'
        response=app.response_class('\ufeff'+buffer.getvalue(),mimetype='text/csv')
        response.headers['Content-Disposition']=f'attachment; filename="{filename}"'
        response.headers['Cache-Control']='no-store'
        return response

    @app.post('/api/uploads/<uid>/reject')
    @require('admin','uploader')
    def reject_upload(uid):
        db().begin_write()
        upload=db().execute('SELECT * FROM uploads WHERE id=?',(uid,)).fetchone()
        if not upload or (g.user['role']!='admin' and upload['user_id']!=g.user['id']):
            raise InvalidData('Upload is unavailable.')
        if upload['state']=='pending':
            db().execute("UPDATE uploads SET state='rejected' WHERE id=?",(uid,))
            log('upload_rejected',json.dumps(dict(id=uid,filename=upload['filename']),separators=(',',':')))
        elif upload['state']=='committed' and g.user['role']=='admin':
            unpublish_upload(uid)
        else:
            raise InvalidData('Only pending files can be rejected; only an admin can unpublish live data.')
        db().commit();return jsonify(ok=True)

    @app.post('/api/uploads/<uid>/archive')
    @require('admin')
    def archive_upload(uid):
        archived=(request.get_json(silent=True) or {}).get('archived')
        if type(archived) is not bool:
            raise InvalidData('Choose whether to archive or unarchive this file.')
        set_upload_archive(uid,archived)
        db().commit();return jsonify(ok=True)

    @app.post('/api/uploads/<uid>/unarchive')
    @require('admin')
    def unarchive_upload(uid):
        set_upload_archive(uid,False)
        db().commit();return jsonify(ok=True)

    def set_upload_archive(uid,archived):
        db().begin_write()
        upload=db().execute('SELECT * FROM uploads WHERE id=?',(uid,)).fetchone()
        if not upload or upload['file_deleted'] or upload['state']=='deleted':
            raise InvalidData('This file has been deleted and cannot be restored.')
        if not archived and upload['state'] in ('pending','rejected'):
            publish_preview(upload,(request.get_json(silent=True) or {}).get('replace') is True)
        elif not archived and upload['state']=='restored':
            rows={(r['day'],r['channel_id']):r for r in resolve_rows(json.loads(upload['rows_json']))}
            revisions=list(db().execute('SELECT * FROM revisions WHERE upload_id=?',(uid,)))
            if not revisions:
                raise InvalidData('This older file has no restorable records. Upload it again to review its data.')
            for revision in revisions:
                key=(revision['day'],revision['channel_id'])
                row=rows.get(key)
                if not row:
                    raise InvalidData('The stored file is incomplete. Upload it again.')
                current=db().execute('SELECT * FROM records WHERE day=? AND channel_id=?',key).fetchone()
                previous=prior_live_record(json.loads(revision['previous']),*key)
                if current and current['upload_id']!=uid and dict(current)!=previous:
                    raise InvalidData('Newer data exists for this file. Upload it again to review and confirm replacements.')
                db().upsert('records',dict(zip(('day','channel_id','views','impressions','ad','other','total','upload_id'),(*key,*row_values(row),uid))))
            db().execute("UPDATE uploads SET state='committed' WHERE id=?",(uid,))
        db().execute('UPDATE uploads SET archived=? WHERE id=?',(int(archived),uid))
        log('upload_archived' if archived else 'upload_unarchived',json.dumps(dict(id=uid,filename=upload['filename']),separators=(',',':')))

    @app.post('/api/uploads/<uid>/delete')
    @require('admin')
    def delete_upload(uid):
        db().begin_write()
        upload=db().execute('SELECT * FROM uploads WHERE id=?',(uid,)).fetchone()
        if not upload:
            raise InvalidData('File not found.')
        if upload['state']=='deleted':
            return jsonify(ok=True)
        try:
            if upload['state']=='committed':
                unpublish_upload(uid)
            log('upload_deleted',json.dumps(dict(id=uid,filename=upload['filename']),separators=(',',':')))
            db().execute("UPDATE uploads SET state='deleted',archived=1,file_deleted=1 WHERE id=?",(uid,))
            db().commit()
        except Exception:
            db().rollback()
            raise
        return jsonify(ok=True)

    @app.post('/api/uploads/<uid>/delete-file')
    @require('admin')
    def delete_upload_file(uid):
        db().begin_write()
        upload=db().execute('SELECT * FROM uploads WHERE id=?',(uid,)).fetchone()
        if not upload or upload['state'] not in ('rejected','restored'):
            raise InvalidData('Reject or unpublish the file before deleting its source. Live data cannot be deleted here.')
        if upload['file_deleted']:
            return jsonify(ok=True)
        log('upload_source_deleted',json.dumps(dict(id=uid,filename=upload['filename']),separators=(',',':')))
        db().execute('UPDATE uploads SET file_deleted=1 WHERE id=?',(uid,))
        db().commit();return jsonify(ok=True)

    def prior_live_record(previous,day,channel_id):
        visited=set()
        while previous:
            source=previous['upload_id']
            if source in visited:
                raise RuntimeError('Upload revision cycle detected.')
            visited.add(source)
            upload=db().execute('SELECT state,file_deleted FROM uploads WHERE id=?',(source,)).fetchone()
            if not upload or (upload['state']=='committed' and not upload['file_deleted']):
                return previous
            revision=db().execute('SELECT previous FROM revisions WHERE upload_id=? AND day=? AND channel_id=?',(source,day,channel_id)).fetchone()
            if not revision:
                raise RuntimeError('Upload revision history is incomplete.')
            previous=json.loads(revision['previous'])
        return None

    def unpublish_upload(uid):
        upload=db().execute('SELECT filename FROM uploads WHERE id=?',(uid,)).fetchone()
        revisions=db().execute('SELECT * FROM revisions WHERE upload_id=?',(uid,)).fetchall()
        restored=0
        for revision in revisions:
            current=db().execute('SELECT upload_id FROM records WHERE day=? AND channel_id=?',(revision['day'],revision['channel_id'])).fetchone()
            if not current or current['upload_id']!=uid:
                continue
            previous=prior_live_record(json.loads(revision['previous']),revision['day'],revision['channel_id'])
            db().execute('DELETE FROM records WHERE day=? AND channel_id=?',(revision['day'],revision['channel_id']))
            if previous:
                db().execute('INSERT INTO records VALUES (?,?,?,?,?,?,?,?)',tuple(previous[key] for key in ('day','channel_id','views','impressions','ad','other','total','upload_id')))
            restored+=1
        db().execute("UPDATE uploads SET state='restored' WHERE id=?",(uid,))
        log('upload_unpublished',json.dumps(dict(id=uid,filename=upload['filename'] if upload else '',affected=restored),separators=(',',':')))

    @app.post('/api/uploads/<uid>/restore')
    @require('admin')
    def restore(uid):
        db().begin_write()
        upload=db().execute("SELECT * FROM uploads WHERE id=? AND state='committed'",(uid,)).fetchone()
        if not upload:
            raise InvalidData('Only a published upload can be rolled back.')
        unpublish_upload(uid)
        db().commit();return jsonify(ok=True)

    @app.get('/api/admin/users')
    @require('admin')
    def users():
        rows=[]
        for user in db().execute('SELECT id,username,role,active,must_change FROM users ORDER BY username'):
            rows.append({**dict(user),'super_admin':bool(db().execute('SELECT 1 FROM super_admin WHERE user_id=?',(user['id'],)).fetchone()),'channels':[r[0] for r in db().execute('SELECT channel_id FROM assignments WHERE user_id=?',(user['id'],))]})
        return jsonify(users=rows,channels=[dict(c) for c in permitted()],archived=[dict(c) for c in db().execute('SELECT c.* FROM channels c JOIN archived_channels a ON a.channel_id=c.id ORDER BY LOWER(c.name)')])

    @app.post('/api/admin/channels/<int:cid>/archive')
    @require('admin')
    def archive_channel(cid):
        archived=(request.get_json(silent=True) or {}).get('archived')
        if type(archived) is not bool:
            raise InvalidData('Specify archive or restore.')
        db().begin_write()
        if not db().execute('SELECT 1 FROM channels WHERE id=?',(cid,)).fetchone():
            raise InvalidData('Channel not found.')
        if archived:
            db().upsert('archived_channels',dict(channel_id=cid),ignore=True)
        else:
            db().execute('DELETE FROM archived_channels WHERE channel_id=?',(cid,))
        log('channel_archived' if archived else 'channel_restored',str(cid))
        db().commit()
        return jsonify(ok=True)

    @app.post('/api/admin/channels')
    @require('admin')
    def channels():
        name=str((request.get_json() or {}).get('name','')).strip()
        if not name or len(name)>120:
            raise InvalidData('Enter a channel name, up to 120 characters.')
        try:
            db().execute('INSERT INTO channels(name) VALUES (?)',(name,))
        except INTEGRITY_ERRORS:
            raise InvalidData('Channel already exists.')
        log('channel_created',name);db().commit();return jsonify(ok=True)

    from account_email import install
    email_address, mail_settings, send_invitation = install(app, db, data, InvalidData, log)

    @app.post('/api/admin/users')
    @require('admin')
    def save_user():
        body=request.get_json() or {}
        username=str(body.get('username','')).strip()
        role=body.get('role')
        password=str(body.get('password',''))
        uid=body.get('id')
        invite = body.get('invite') is True
        if invite:
            if uid:
                raise InvalidData('Invitations are for new accounts. Existing email users can use Forgot password.')
            username = email_address(username)
            mail_settings()
            password = secrets.token_urlsafe(48)
        if uid is not None and (type(uid)!=int or uid<=0):
            raise InvalidData('Invalid user ID.')
        if role not in {'admin','uploader','viewer'} or not username or len(username)>80:
            raise InvalidData('Enter a username and valid role.')
        if (not uid or password) and not 12<=len(password)<=256:
            raise InvalidData('Temporary password must be 12 to 256 characters.')
        ids=body.get('channels',[])
        if not isinstance(ids,list) or any(type(x)!=int for x in ids):
            raise InvalidData('Invalid channel assignments.')
        all_ids={r[0] for r in db().execute('SELECT id FROM channels')}
        if not set(ids)<=all_ids:
            raise InvalidData('Unknown channel assignment.')
        active=int(body.get('active',True) is True)
        if uid==g.user['id'] and (role!='admin' or not active):
            raise InvalidData('You cannot remove your own admin access.')
        try:
            db().begin_write()
            owner=db().execute('SELECT user_id FROM super_admin WHERE singleton=1').fetchone()
            existing=db().execute('SELECT username,role,active FROM users WHERE id=?',(uid,)).fetchone() if uid else None
            if owner and uid==owner['user_id']:
                raise InvalidData('The Super Admin account is protected. Use Change password for your own password.')
            if not owner or owner['user_id']!=g.user['id']:
                if role=='admin' or (existing and existing['role']=='admin'):
                    return jsonify(error='Only the Super Admin can create or modify admin accounts.'),403
            if uid:
                if not db().execute('SELECT 1 FROM users WHERE id=?',(uid,)).fetchone():
                    raise InvalidData('User not found.')
                before={**dict(existing),'channels':[row[0] for row in db().execute('SELECT channel_id FROM assignments WHERE user_id=? ORDER BY channel_id',(uid,))]}
                email_account=db().execute('SELECT email FROM email_accounts WHERE user_id=?',(uid,)).fetchone()
                if email_account and username.casefold()!=email_account['email'].casefold():
                    raise InvalidData('Email login cannot be renamed. Disable this account and invite the new email separately.')
                db().execute('UPDATE users SET username=?,role=?,active=? WHERE id=?',(username,role,active,uid))
                if password:
                    db().execute('UPDATE users SET password=?,must_change=1 WHERE id=?',(generate_password_hash(password),uid))
                if password or not active or existing['role'] != role:
                    db().execute('DELETE FROM sessions WHERE user_id=? AND token<>?',(uid,g.session['token']))
                    db().execute('DELETE FROM email_tokens WHERE user_id=?',(uid,))
            else:
                before=None
                uid=db().execute('INSERT INTO users(username,password,role,active,must_change) VALUES (?,?,?,?,1)',(username,generate_password_hash(password),role,active)).lastrowid
            db().execute('DELETE FROM assignments WHERE user_id=?',(uid,))
            db().executemany('INSERT INTO assignments VALUES (?,?)',[(uid,c) for c in set(ids)])
            if invite:
                db().execute('INSERT INTO email_accounts(user_id,email) VALUES (?,?)',(uid,username))
            log('user_saved',json.dumps(dict(id=uid,username=username,role=role,active=bool(active),channels=sorted(set(ids)),before=before,password_reset=bool(password and before),invited=invite),separators=(',',':')))
            db().commit()
        except INTEGRITY_ERRORS:
            raise InvalidData('Username already exists.')
        if invite:
            send_invitation(uid)
        return jsonify(ok=True)

    return app


def bootstrap(data):
    app=create_app(data)
    with closing(app.extensions['database'].connect()) as con, con:
        con.begin_write()
        if not con.execute('SELECT 1 FROM users LIMIT 1').fetchone():
            password=secrets.token_urlsafe(18)
            uid=con.execute("INSERT INTO users(username,password,role,must_change) VALUES (?,?,'admin',1)",('admin',generate_password_hash(password))).lastrowid
            con.execute('INSERT INTO super_admin VALUES (1,?)',(uid,))
            (data/'initial_admin.txt').write_text('Username: admin\nTemporary password: '+password+'\nChange at first login. Keep this file private and delete after changing.\n',encoding='utf-8')
    return app


def backup_database(data, backup_root=None):
    settings = Settings.from_env(data)
    if settings.db_url.get_backend_name() == 'mysql':
        raise ValueError('Use MySQL backup tooling and back up UPLOAD_DIR; SQLite backup is unavailable in MySQL mode.')
    data = Path(data).resolve()
    if not (data / 'revenuelive.db').is_file():
        raise FileNotFoundError('No RevenueLive database found to back up.')
    backup_root = Path(backup_root or os.environ.get('REVENUE_BACKUP_DIR', data.parent / 'RevenueLive-backups')).resolve()
    if backup_root == data or backup_root.is_relative_to(data):
        raise ValueError('Keep backups outside the RevenueLive data directory.')
    target = backup_root / dt.datetime.now(dt.timezone.utc).strftime('%Y%m%d_%H%M%S_%f')
    target.mkdir(parents=True)
    with closing(sqlite3.connect(data / 'revenuelive.db')) as source, closing(sqlite3.connect(target / 'revenuelive.db')) as dest:
        source.backup(dest)
    (target / 'BACKUP_COMPLETE').write_text('SQLite database copied successfully. Upload rows are stored in the database.\n', encoding='ascii')
    return target


if __name__=='__main__':
    load_environment()
    parser=argparse.ArgumentParser()
    parser.add_argument('--host',default='127.0.0.1')
    parser.add_argument('--port',type=int,default=8820)
    parser.add_argument('--backup',action='store_true',help='Back up the real database and uploads, then exit.')
    args=parser.parse_args()
    settings = Settings.from_env()
    data = settings.data_dir
    from runtime_logging import configure_logging
    configure_logging(data)
    if args.backup:
        print('Backup:', backup_database(data), flush=True)
        raise SystemExit(0)
    if settings.app_url and args.host not in {'127.0.0.1','localhost','::1'}:
        parser.error('HTTPS proxy mode requires a loopback backend listener.')
    app=create_app() if settings.db_url.get_backend_name() in {'mysql','mariadb'} else bootstrap(data)
    from waitress import create_server
    options=dict(host=args.host,port=args.port,threads=4,clear_untrusted_proxy_headers=True)
    trusted_proxy=os.environ.get('REVENUE_TRUSTED_PROXY','').strip()
    if trusted_proxy:
        if not settings.app_url:
            parser.error('Trusted proxy headers require HTTPS proxy mode.')
        options.update(trusted_proxy=trusted_proxy,trusted_proxy_headers={'x-forwarded-for'})
    server=create_server(app,**options)
    stop=data/f'stop_{args.port}'
    stop.unlink(missing_ok=True)
    def monitor():
        while not stop.exists():
            time.sleep(1)
        server.close()
        os._exit(0)
    threading.Thread(target=monitor,daemon=True).start()
    logging.getLogger(__name__).info('RevenueLive listening on %s:%s',args.host,args.port)
    server.run()
